"""在服务器上读客户文件：PDF 先读文字层，读不到再文字识别（OCR）；图片走 OCR；Excel 按栏名；DXF 读文字和尺寸标注；
DWG 先用免费的 ODA File Converter 转成 DXF 副本。原件只读，截图和转换结果都另存在 derived 目录。
PDF 和图片每页读完后再做一次尺寸线配对（单独的「6400mm」按位置当作屏宽 / 屏高候选，见 dimensions.py）。"""
import math
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from .dimensions import image_rects, pair
from .fields import LABELS, find_in_text, parse_value, to_mm

MAX_PAGES = 30  # 文字层最多读前 30 页
MAX_OCR_PAGES = 5  # 扫描件最多识别前 5 页（CPU 上较慢）
MAX_SIDE_PX = 3000  # 送去识别的图片最长边

_ocr = None


def ocr_lines(img) -> list[tuple[str, tuple[float, float, float, float]]]:
    """图片里的文字行：(文字, 外框)。识别模型第一次用时才加载。"""
    global _ocr
    if _ocr is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr = RapidOCR()
    import numpy as np
    result, _ = _ocr(np.array(img.convert("RGB")))
    lines = []
    for box, text, _score in result or []:
        xs, ys = [p[0] for p in box], [p[1] for p in box]
        lines.append((text, (min(xs), min(ys), max(xs), max(ys))))
    return lines


def _shrink(img):
    scale = MAX_SIDE_PX / max(img.size)
    return img.resize((int(img.width * scale), int(img.height * scale))) if scale < 1 else img


class Collector:
    """收集一个文件里认出的字段，同字段同值只留第一条，截图存到 out_dir。"""

    def __init__(self, out_dir: Path, rel_dir: str):
        self.out_dir, self.rel_dir, self.fields, self.text = out_dir, rel_dir, [], []

    def line(self, text: str, ocr: bool, page=None, location="", crop=None):
        self.text.append(text)
        for field, value, conf, snippet in find_in_text(text, ocr):
            self.add(field, value, conf, page, location, snippet, crop)

    def dimensions(self, lines, rects, ocr: bool, page, location: str, crop):
        """尺寸线配对：OCR 读到的给 low，PDF 文字层读到的给 medium，都要在确认关逐项确认。"""
        for field, value, text, box in pair(lines, rects):
            self.add(field, value, "low" if ocr else "medium", page, f"{location} · 尺寸线", text,
                     lambda b=box: crop(b))

    def add(self, field, value, conf, page, location, snippet, crop=None):
        if any(f["field"] == field and f["value"] == value for f in self.fields):
            return
        path = None
        if crop is not None:
            name = f"crop-{len(self.fields) + 1}.png"
            crop().save(self.out_dir / name)
            path = f"{self.rel_dir}/{name}"
        self.fields.append({"field": field, "value": value, "confidence": conf, "page": page,
                            "location": location, "snippet": snippet[:300], "crop": path})


def read_pdf(path: Path, c: Collector) -> str:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        pages = pdf.pages[:MAX_PAGES]
        if any((p.extract_text() or "").strip() for p in pages):
            for no, page in enumerate(pages, 1):
                def crop(b, p=page):
                    b = (max(b[0] - 6, 0), max(b[1] - 6, 0), min(b[2] + 6, p.width), min(b[3] + 6, p.height))
                    return p.crop(b).to_image(resolution=150).original
                lines = [(line["text"], (line["x0"], line["top"], line["x1"], line["bottom"]))
                         for line in page.extract_text_lines()]
                for text, box in lines:
                    c.line(text, ocr=False, page=no, location=f"p.{no}", crop=lambda b=box, f=crop: f(b))
                rects = [(r["x0"], r["top"], r["x1"], r["bottom"]) for r in page.rects]
                c.dimensions(lines, rects, ocr=False, page=no, location=f"p.{no}", crop=crop)
            return "text_layer"
        for no, page in enumerate(pdf.pages[:MAX_OCR_PAGES], 1):  # 扫描件：整页渲染后识别
            dpi = min(200, MAX_SIDE_PX / (max(page.width, page.height) / 72))
            img = page.to_image(resolution=dpi).original
            _ocr_image(img, c, page=no)
        return "ocr"


def _ocr_image(img, c: Collector, page=None):
    img = _shrink(img)
    location = f"p.{page} · 文字识别" if page else "文字识别"

    def crop(b):
        return img.crop((max(b[0] - 8, 0), max(b[1] - 8, 0), min(b[2] + 8, img.width), min(b[3] + 8, img.height)))

    lines = ocr_lines(img)
    for text, box in lines:
        c.line(text, ocr=True, page=page, location=location, crop=lambda b=box: crop(b))
    c.dimensions(lines, image_rects(img), ocr=True, page=page, location=location, crop=crop)


def read_image(path: Path, c: Collector) -> str:
    from PIL import Image, ImageOps

    with Image.open(path) as img:
        _ocr_image(ImageOps.exif_transpose(img).convert("RGB"), c)
    return "ocr"


def read_excel(path: Path, c: Collector) -> str:
    """按栏名匹配：栏名在左、值在右（或栏名在上、值在下）。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=False)
    for ws in wb.worksheets:
        for row in ws.iter_rows(max_row=200, max_col=30):
            for cell in row:
                if not isinstance(cell.value, str) or len(cell.value) > 40:
                    continue
                c.text.append(cell.value)
                field = next((f for f, pat in LABELS.items() if re.search(pat, cell.value, re.I)), None)
                if field is None:
                    continue
                right = [ws.cell(cell.row, col) for col in range(cell.column + 1, cell.column + 4)]
                below = [ws.cell(r, cell.column) for r in range(cell.row + 1, cell.row + 2)]
                for target in right + below:
                    if target.value not in (None, ""):
                        value = parse_value(field, target.value)
                        if value is not None:
                            c.add(field, value, "high", None, f"{ws.title}!{target.coordinate}",
                                  f"{cell.value}: {target.value}")
                        break
    return "excel"


DXF_UNIT_MM = {0: 1, 1: 25.4, 4: 1, 5: 10, 6: 1000}  # $INSUNITS：未设定按 mm 算


def read_dxf(path: Path, c: Collector) -> str:
    import ezdxf

    doc = ezdxf.readfile(path)
    msp = doc.modelspace()
    for e in msp.query("TEXT MTEXT"):
        text = e.plain_text() if e.dxftype() == "MTEXT" else e.dxf.text
        x, y = e.dxf.insert.x, e.dxf.insert.y
        c.line(text, ocr=False, location=f"{e.dxf.layer} · ({x:.0f}, {y:.0f})")
    # 尺寸标注（DIMENSION）：取实测值，横向的算屏宽候选、纵向的算屏高候选，各取最大的一个
    scale = DXF_UNIT_MM.get(doc.header.get("$INSUNITS", 0))
    found = {"width_mm": [], "height_mm": []}
    for d in msp.query("DIMENSION") if scale else []:
        side = _dim_side(d)
        value = side and to_mm(str(round(d.get_measurement() * scale)), "mm")
        if value:
            found[side].append((value, d.dxf.layer))
    for field, items in found.items():
        if items:
            value, layer = max(items)
            c.add(field, value, "medium", None, f"{layer} · 尺寸标注", f"DIMENSION {value}")
    return "dxf"


def _dim_side(d) -> str | None:
    """线性标注看角度，对齐标注看两个定义点的方向；角度、半径等其它标注不要。"""
    kind = d.dimtype & 7
    if kind == 0:
        angle = d.dxf.get("angle", 0) % 180
    elif kind == 1:
        p1, p2 = d.dxf.defpoint2, d.dxf.defpoint3
        angle = abs(math.degrees(math.atan2(p2.y - p1.y, p2.x - p1.x))) % 180
    else:
        return None
    return "width_mm" if min(angle, 180 - angle) < 10 else "height_mm" if abs(angle - 90) < 10 else None


def convert_dwg(path: Path, out_dir: Path, oda: str) -> Path:
    """ODA File Converter 只处理整个文件夹：把 DWG 复制到临时文件夹再转换，DXF 副本存到 out_dir。"""
    with tempfile.TemporaryDirectory() as src, tempfile.TemporaryDirectory() as dst:
        shutil.copy(path, Path(src) / "source.dwg")
        subprocess.run([oda, src, dst, "ACAD2018", "DXF", "0", "1", "*.DWG"], check=True, timeout=180,
                       capture_output=True)
        dxf = next(Path(dst).glob("*.dxf"))
        target = out_dir / "converted.dxf"
        shutil.copy(dxf, target)
        return target
