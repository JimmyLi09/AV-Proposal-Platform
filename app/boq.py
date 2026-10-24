"""客户版 BOQ（CR-6）：按模板把规则结果写成汇总行，不含价格。

数量来自规则（或模板里的固定数，例如 1 lot）；人改过的数量标「手动」并写原因，规则结果本身不变。
箱体、电源、接收卡、线缆根数等留在内部明细（规则引擎的 boq 行），不导出。
"""
import re
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

NAME = re.compile(r"\{(\w+)\}")


def _text(v) -> str:
    if v is None or v in ("", []):
        return "—"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{Decimal(str(v)).quantize(Decimal('0.01'), ROUND_HALF_UP).normalize():f}"
    return str(v)


def fill(text: str, ctx: dict) -> str:
    return NAME.sub(lambda m: _text(ctx.get(m[1])), text)


def _qty(spec, ctx: dict, unit: str):
    v = spec if isinstance(spec, (int, float)) else ctx.get(spec)
    if not isinstance(v, (int, float)) or isinstance(v, bool):
        return None
    v = round(float(v), 2)
    return int(v) if v.is_integer() and unit != "sqm" else v


def build(template: dict, ctx: dict, edits: dict | None, manual: set = frozenset()) -> dict:
    """ctx：规则结果、输入和方案名；edits：这个 Rev 里人改过的内容（rev.boq_edits）；
    manual：推荐页手动覆盖过的结果名（CR-5），数量取自这些结果的行标 rule_manual。"""
    edits = edits or {}
    over = edits.get("rows") or {}
    rows = []
    for r in template["boq"]["rows"]:
        if "when" in r and ctx.get(r["when"]) in (0, False):
            continue
        title, specs = fill(r["title"], ctx), [fill(s, ctx) for s in r["specs"]]
        rule_qty = _qty(r["qty"], ctx, r["unit"])
        o = over.get(r["key"], {})
        rows.append({"key": r["key"], "title": o.get("title", title), "title_default": title,
                     "specs": o.get("specs", specs), "specs_default": specs,
                     "qty": o.get("qty", rule_qty), "rule_qty": rule_qty, "unit": r["unit"],
                     "note": fill(r.get("note", ""), ctx), "source": r.get("source", ""),
                     "manual": "qty" in o, "reason": o.get("reason"), "by": o.get("by"), "at": o.get("at"),
                     "added": False, "rule_manual": r["qty"] in manual})
    for a in edits.get("added") or []:
        rows.append({"key": a["key"], "title": a["title"], "title_default": None, "specs": a["specs"],
                     "specs_default": None, "qty": a["qty"], "rule_qty": None, "unit": a["unit"], "note": "",
                     "source": "", "manual": True, "reason": None, "by": a["by"], "at": a["at"], "added": True,
                     "rule_manual": False})
    for n, row in enumerate(rows, 1):
        row["no"] = n
    group_default = fill(template["boq"]["group"], ctx)
    return {"group": edits.get("group") or group_default, "group_default": group_default, "rows": rows,
            "manual": [{k: row[k] for k in ("no", "key", "title", "rule_qty", "qty", "unit", "reason", "by", "at",
                                            "added")} for row in rows if row["manual"]]}


def build_multi(template: dict, faces: list[tuple[str, str, dict]], shared_title: str, shared_ctx: dict,
                pooled: tuple = ()) -> dict:
    """多面屏（CR-7）：数量取自规则结果的行每面一组（faces：面编号、组标题、这一面的值）；固定数量的行（例如 1 set 的电脑）
    只出一次，放在最后的共用组。pooled：整屏合计的值（同型号备品），用到它们的规格行从每面拿掉，在共用组单独出一行。
    目前只读，不能在页面上改。"""
    def part(rows, ctx):
        return build({**template, "boq": {**template["boq"], "rows": rows}}, ctx, None)["rows"] if rows else []

    is_pooled = lambda s: bool(set(NAME.findall(s)) & set(pooled))  # noqa: E731
    by_rule = [r | {"specs": [s for s in r["specs"] if not is_pooled(s)]}
               for r in template["boq"]["rows"] if isinstance(r["qty"], str)]
    fixed = [r for r in template["boq"]["rows"] if not isinstance(r["qty"], str)]
    fixed += [{"key": f"pooled:{r['key']}", "title": s, "specs": [], "qty": 1, "unit": "lot"}
              for r in template["boq"]["rows"] if isinstance(r["qty"], str) for s in r["specs"] if is_pooled(s)]
    rows = []
    for fid, title, ctx in faces:
        rows += [r | {"key": f"{fid}:{r['key']}", "group": title} for r in part(by_rule, ctx)]
    rows += [r | {"group": shared_title} for r in part(fixed, shared_ctx)]
    for n, row in enumerate(rows, 1):
        row["no"] = n
    group = fill(template["boq"]["group"], shared_ctx)
    return {"group": group, "group_default": group, "rows": rows, "manual": [], "readonly": True}


def to_xlsx(boq: dict, meta: dict) -> bytes:
    """客户版 BOQ Excel：No.｜Description｜Qty｜Unit。meta：标题、副标题、版本说明（Rev、规则、产品库、模板）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    ws.append([meta["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([meta["subtitle"]])
    ws.append([meta["stamp"]])
    ws.append([])
    ws.append(["No.", "Description", "Qty", "Unit"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="15203A")
    ws.append(["", boq["group"]])
    ws.cell(ws.max_row, 2).font = Font(bold=True)
    group = None
    for r in boq["rows"]:
        if r.get("group") and r["group"] != group:  # 多面屏按面分组
            group = r["group"]
            ws.append(["", group])
            ws.cell(ws.max_row, 2).font = Font(bold=True, italic=True)
        ws.append([r["no"], r["title"], r["qty"], r["unit"]])
        ws.cell(ws.max_row, 2).font = Font(bold=True)
        for s in r["specs"]:
            ws.append(["", f"- {s}"])
        if r["note"]:
            ws.append(["", f"*{r['note']}"])
            ws.cell(ws.max_row, 2).font = Font(italic=True)
    for col, width in zip("ABCD", (6, 90, 10, 8)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=5):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    p = wb.properties
    p.title, p.subject, p.keywords, p.description = meta["title"], meta["subtitle"], meta["stamp"], meta["stamp"]
    p.creator = "AV Proposal Platform"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
