"""CR-7：投影方案和多面屏方案在页面接口上的行为（新建、录入、计算、客户版 BOQ、确认关、英文界面、升级）。"""
import copy
import io
import json
import re
from pathlib import Path

import sqlalchemy as sa
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.db import MIGRATIONS, make_engine, migrate
from app.main import create_app
from app.rules.engine import load_seed

SAMPLES = Path(__file__).parent / "fixtures" / "samples"
CASES = json.loads((Path(__file__).parent / "fixtures" / "regression_cases.json").read_text(encoding="utf-8"))
CJK = re.compile(r"[　-〿一-鿿＀-￯]")
LED_RULES, LED_SAMPLE = load_seed("LED")
_, PJ_SAMPLE = load_seed("PJ")


def start(tmp_path, line):
    c = TestClient(create_app(tmp_path))
    c.__enter__()
    user = c.post("/api/users", json={"name": "Jimmy", "role": "PM"}).json()
    c.post("/api/session", json={"user_id": user["id"]})
    c.project = c.post("/api/projects", json={"name": "CR-7", "client": "Audax"}).json()["id"]
    c.proposal = c.post(f"/api/projects/{c.project}/proposals", json={"name": "Wall-1", "business_line": line}).json()["id"]
    return c


def put_inputs(c, inputs):
    r = c.put(f"/api/proposals/{c.proposal}/inputs", json={"inputs": inputs})
    assert r.status_code == 200, r.text
    return r.json()


def test_projection_proposal_end_to_end(tmp_path):
    c = start(tmp_path, "PJ")
    v = put_inputs(c, PJ_SAMPLE["input"])
    assert v["rev"]["status"] == "calculated" and v["rules_version"] == "v1.5" and v["shapes"][0] == "single"
    rows = {r["key"]: r for r in v["boq"]["rows"]}
    assert rows["projector"]["title"] == "Panasonic PT-BHZ611C laser projector" and rows["projector"]["qty"] == 5
    assert rows["lens"]["qty"] == 5 and rows["blender"]["qty"] == 1 and rows["transmission"]["qty"] == 5
    assert rows["cabling"]["specs"] == ["Video 5 + 1 spare (HDBaseT)", "Control CAT6 5 + 1 spare"]
    assert v["result"]["diagram"]["unmatched"] == [] and v["drawing"] is None
    xlsx = load_workbook(io.BytesIO(c.get(f"/api/proposals/{c.proposal}/boq.xlsx").content))
    assert any(row[1] == "Content production" for row in xlsx.active.iter_rows(values_only=True))
    en = c.get(f"/api/proposals/{c.proposal}", headers={"X-Lang": "en"}).json()
    leaks = [s for s in json.dumps(en["result"]["rules"], ensure_ascii=False).split('"') if CJK.search(s)]
    assert leaks == []
    for path in ("/api/rules/PJ", "/api/products/PJ", "/api/templates/PJ"):
        assert c.get(path).status_code == 200


def test_conflicting_file_values_wait_for_confirmation_on_a_projection_proposal(tmp_path):
    """Phoenix Sea「正文 4 台、图注 6 台」：同一个字段几份文件读到的不一样，进确认关，确认前不能计算。"""
    c = start(tmp_path, "PJ")
    files = [("files", (n, (SAMPLES / n).read_bytes())) for n in ("Lobby_Elevation_RevB.pdf", "Lobby_site_photo.jpg")]
    assert c.post(f"/api/projects/{c.project}/files", files=files).status_code == 201
    v = c.post(f"/api/proposals/{c.proposal}/fill-from-files").json()
    assert "wall_width_mm" in v["pending"] and v["rev"]["input_meta"]["wall_width_mm"]["conflict"]
    assert v["rev"]["status"] == "confirming" and v["result"] is None


def gmc_inputs():
    i = CASES["led_gmc_u_shape"]["input"]
    faces = [{k: f[k] for k in ("id", "name", "orientation", "width_mm", "height_mm")} for f in i["faces"]]
    return {k: x for k, x in LED_SAMPLE["input"].items() if k not in ("width_mm", "height_mm")} | {
        "shape": "U", "faces": faces, "corners": i["corners"]}


def test_multiface_led_proposal(tmp_path):
    c = start(tmp_path, "LED")
    v = put_inputs(c, gmc_inputs())
    # 示例箱体 640 宽：B 面 8000 拼不出（F3 阻断），用到箱体数的 F8 缺数，方案停在草稿；其它数量照样按面算出
    assert v["rev"]["status"] == "draft" and "B · F8" in next(r for r in v["result"]["rules"] if r["id"] == "S3")["missing"]
    assert [f["id"] for f in v["result"]["faces"]] == ["A", "B", "C"] and v["result"]["multi"]["totals"]["circuits"] == 17
    boq = v["boq"]
    assert boq["readonly"] and [r["group"] for r in boq["rows"]][:1] == ["A · Left"]
    assert {r["group"] for r in boq["rows"]} == {"A · Left", "B · Main", "C · Right", "共用"}
    # 同型号备品按整屏合计（S3）：不在每面重复，只在共用组出一行；共用组的线缆备用数按面加总
    shared = [r for r in boq["rows"] if r["group"] == "共用"]
    assert shared[-1]["title"] == "Spares: 80 modules, 3 power supply, 3 receiving card"
    assert "Power cables 17 + 3 spare, data cables CAT6 24 + 3 spare" in shared[1]["specs"]
    assert not any(s.startswith("Spares") for r in boq["rows"] if r["group"] != "共用" for s in r["specs"])
    assert c.put(f"/api/proposals/{c.proposal}/boq", json={"rows": {}}).status_code == 409
    rows = list(load_workbook(io.BytesIO(c.get(f"/api/proposals/{c.proposal}/boq.xlsx").content)).active.iter_rows(values_only=True))
    assert (None, "B · Main", None, None) in rows
    assert v["result"]["diagram"] is None and v["drawing"] is None  # 示意图、图纸目前只做平面单屏
    en = c.get(f"/api/proposals/{c.proposal}", headers={"X-Lang": "en"}).json()["result"]
    leaks = [s for s in json.dumps([en["rules"], en["multi"]], ensure_ascii=False).split('"') if CJK.search(s)]
    assert leaks == []

    conflict = gmc_inputs()
    conflict["faces"][0] |= {"width_mm": None, "width_sources": [{"src": "平面图", "kind": "plan", "value": 5600},
                                                                 {"src": "LED 立面图", "kind": "elevation", "value": 5760}]}
    v = put_inputs(c, conflict)
    assert v["rev"]["status"] == "confirming" and v["result"]["multi"]["conflicts"][0]["face"] == "A"


def test_upgrade_adds_the_s5_output_limit_as_a_new_rule_version(tmp_path):
    engine = make_engine(tmp_path / "avp.db")
    cfg = Config()
    cfg.set_main_option("script_location", str(MIGRATIONS))
    with engine.begin() as conn:
        cfg.attributes["connection"] = conn
        command.upgrade(cfg, "0009")
    old = copy.deepcopy(LED_RULES)
    old["params"] = [p for p in old["params"] if p["name"] != "output_max_px_w"]
    sets = sa.table("rule_set_versions", sa.column("business_line"), sa.column("version"),
                    sa.column("content", sa.JSON), sa.column("reason"), sa.column("created_at", sa.DateTime))
    with engine.begin() as conn:
        conn.execute(sets.insert().values(business_line="LED", version="v1.4", content=old, reason="试用",
                                          created_at=sa.func.now()))
    migrate(engine)
    with engine.connect() as conn:
        rows = conn.execute(sa.text("select version, content, reason from rule_set_versions order by id")).all()
    assert [r.version for r in rows] == ["v1.4", "v1.5"] and "CR-7" in rows[1].reason
    new = json.loads(rows[1].content)
    assert new["rules"] == old["rules"] and new["params"][-1]["name"] == "output_max_px_w"
