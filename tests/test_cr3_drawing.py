"""CR-3 验收：拼接与线路图。143-Royal Plaza：35 个箱体；回路数 = F6，每个回路不超过回路容量；网线数 = F9 + 1 备用；
尺寸线加起来 4480 / 2560。DXF 图层齐全、尺寸可量、1:1；手动调整后超载要标红、不能批准；6400 × 3520 也能出图。"""
import io
import json
from pathlib import Path

import ezdxf
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.drawing import build
from app.main import create_app
from app.products.library import load_seed as load_library
from app.rules.engine import calculate, load_seed

RULES, SAMPLE = load_seed("LED")
PARAMS = {p["name"]: p["value"] for p in RULES["params"]}
TRIAL = SAMPLE["input"] | {"width_mm": 6400, "height_mm": 3520}


def drawing(inputs, edits=None):
    out = calculate(RULES, load_library("LED"), inputs)
    values = {ln["name"]: ln["value"] for r in out["rules"] for ln in r["lines"]}
    return build(values, inputs, PARAMS, edits), values


def test_royal_plaza_drawing_matches_the_rules():
    d, v = drawing(SAMPLE["input"])
    assert len(d["cabinets"]) == v["cabinets"] == 35
    assert {k: s["count"] for k, s in d["sizes"].items()} == {"A": 28, "B": 7}
    assert d["rows"] == [480, 480, 480, 480, 640] and d["cols"] == 7  # 主尺寸的行在下面
    assert sum(d["rows"]) == 2560 and d["cols"] * d["col_w"] == 4480
    assert [c["id"] for c in d["cabinets"][:2]] == ["R1C1", "R1C2"] and d["cabinets"][0]["y"] == 0  # 从左下角编号
    assert len(d["circuits"]) == v["circuits"] == 3 and all(c["watts"] <= 2500 for c in d["circuits"])
    assert sum(c["watts"] for c in d["circuits"]) == pytest.approx(v["power_kw"] * 1000, abs=2)
    assert len(d["cables"]) == v["data_runs"] == 4 and d["spare_cables"] == 1 and d["spare_circuits"] == 1
    assert all(c["px"] <= 560_000 for c in d["cables"]) and sum(c["px"] for c in d["cables"]) == v["total_pixels"]
    assert d["cables"][0]["path"][:6] == ["R1C1", "R2C1", "R3C1", "R4C1", "R5C1", "R5C2"]  # 左边起，按列蛇形
    assert d["checks"] == [] and not d["blocked"]


def test_trial_6400_x_3520_is_drawn_with_7_plus_1_data_cables():
    d, v = drawing(TRIAL)
    assert len(d["cabinets"]) == 70 and len(d["cables"]) == 7 == v["data_runs"] and len(d["circuits"]) == v["circuits"] == 5
    assert d["checks"] == []


def test_manual_regroup_overloads_and_blocks():
    d, _ = drawing(SAMPLE["input"], {"power": {"R1C7": 1, "R2C7": 1}, "data": {"R1C3": 1}})
    p1, d1 = d["circuits"][0], d["cables"][0]
    assert p1["over"] and p1["watts"] > 2500 and d1["over"] and d1["px"] > 560_000
    assert d["blocked"] and [c["check"] for c in d["checks"]] == ["BLOCK", "BLOCK"]
    assert d["manual"] == {"power": 2, "data": 1}
    assert next(c for c in d["cabinets"] if c["id"] == "R1C7")["manual"]


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path)) as c:
        user = c.post("/api/users", json={"name": "Jimmy", "role": "PM"}).json()
        c.post("/api/session", json={"user_id": user["id"]})
        pid = c.post("/api/projects", json={"name": "143-Royal Plaza LED", "client": "Royal Plaza"}).json()["id"]
        c.proposal = c.post(f"/api/projects/{pid}/proposals", json={"name": "LED-1 Lobby", "business_line": "LED"}).json()["id"]
        c.put(f"/api/proposals/{c.proposal}/inputs", json={"inputs": SAMPLE["input"]})
        yield c


def get(client, path):
    r = client.get(f"/api/proposals/{client.proposal}/{path}")
    assert r.status_code == 200, r.text
    return r


def test_drawing_in_view_with_calculation_basis(client):
    d = client.get(f"/api/proposals/{client.proposal}").json()["drawing"]
    assert len(d["cabinets"]) == 35 and [b["rule"] for b in d["basis"]] == ["F3", "F4", "F5", "F6", "F7", "F9"]
    assert d["basis"][-1]["formula"] == "ceil(1,835,008 / 560,000) = 4；4 + 1 = 5"


def test_manual_adjustment_is_saved_with_the_rev_and_reset(client):
    url = f"/api/proposals/{client.proposal}/drawing"
    d = client.put(url, json={"power": {"R1C7": 1, "R2C7": 1}, "data": {}}).json()["drawing"]
    assert d["blocked"] and d["manual"]["power"] == 2
    assert client.get(f"/api/proposals/{client.proposal}").json()["drawing"]["blocked"]  # 随 Rev 保存
    assert client.put(url, json={"power": {"R1C7": 9}}).status_code == 422
    assert client.put(url, json={"power": {"R9C9": 1}}).status_code == 422
    d = client.put(url, json={}).json()["drawing"]  # 恢复自动
    assert not d["blocked"] and d["manual"] == {"power": 0, "data": 0}


def test_dxf_layers_units_and_measurable_dimensions(client):
    r = get(client, "drawing.dxf")
    assert "Drawing_Rev0_DRAFT.dxf" in r.headers["content-disposition"]
    doc = ezdxf.read(io.StringIO(r.content.decode("utf-8")))
    assert doc.header["$INSUNITS"] == 4  # mm，1:1
    layers = {l.dxf.name for l in doc.layers}
    assert {"AVP-CABINET", "AVP-CABINET-ID", "AVP-POWER", "AVP-DATA", "AVP-DIM", "AVP-NOTE"} <= layers
    msp = doc.modelspace()
    assert len(msp.query('LWPOLYLINE[layer=="AVP-CABINET"]')) == 35
    dims = [round(x.get_measurement()) for x in msp.query("DIMENSION")]
    assert dims[:7] == [640] * 7 and dims[7] == 4480 and dims[8:13] == [480, 480, 480, 480, 640] and dims[13] == 2560
    assert all(x.dxf.layer == "AVP-DIM" for x in msp.query("DIMENSION"))
    assert len(msp.query('LWPOLYLINE[layer=="AVP-DATA"]')) == 4
    assert any("DRAFT" in t.dxf.text for t in msp.query('TEXT[layer=="AVP-NOTE"]'))


def test_svg_excel_and_pdf_exports(client):
    svg = get(client, "drawing.svg").text
    assert svg.startswith("<svg") and svg.count('data-cab="') == 35 and "DRAFT" in svg and "4480 = 640 × 7" in svg
    assert "layer-note" not in get(client, "drawing.svg?layers=id,size").text  # 图层开关
    wb = load_workbook(io.BytesIO(get(client, "cabinets.xlsx").content))
    rows = list(wb["Cabinets"].iter_rows(values_only=True))
    assert rows[3][:6] == ("ID", "Size 尺寸", "Circuit 回路", "Data cable 网线", "Row 行", "Column 列")
    assert len(rows) == 4 + 35 and rows[4][:6] == ("R1C1", "A 640 × 480", 1, 1, 1, 1)
    assert wb.properties.keywords == "Rev 0 · Rules v1.3 · Product library v1.0 · DRAFT"
    for lang, name in (("zh", "%E6%8A%80%E6%9C%AF%E6%96%B9%E6%A1%88"), ("en", "Technical")):
        r = get(client, f"drawing.pdf?lang={lang}")
        assert r.content[:5] == b"%PDF-" and name in r.headers["content-disposition"]
        meta = PdfReader(io.BytesIO(r.content)).metadata
        assert meta["/Keywords"] == "Rev 0 · Rules v1.3 · Product library v1.0 · DRAFT"


def test_dwg_needs_oda(client):
    r = client.get(f"/api/proposals/{client.proposal}/drawing.dwg")
    assert r.status_code == 409 and "ODA" in r.json()["detail"]
    assert client.get("/api/capabilities").json() == {"dwg": False}


def test_dwg_is_converted_from_the_dxf_when_oda_is_installed(tmp_path):
    """测试用：假装 ODA 把文件夹里的 DXF 转成 DWG（只是改个名）。"""
    import os
    import stat
    if os.name == "nt":
        oda = tmp_path / "fake_oda.cmd"
        oda.write_text('@copy "%~1\\drawing.dxf" "%~2\\drawing.dwg" >nul\n')
    else:
        oda = tmp_path / "fake_oda.sh"
        oda.write_text('#!/bin/sh\ncp "$1/drawing.dxf" "$2/drawing.dwg"\n')
        oda.chmod(oda.stat().st_mode | stat.S_IEXEC)
    with TestClient(create_app(tmp_path / "data")) as c:
        c.app.state.oda = str(oda)
        user = c.post("/api/users", json={"name": "Jimmy", "role": "PM"}).json()
        c.post("/api/session", json={"user_id": user["id"]})
        pid = c.post("/api/projects", json={"name": "P", "client": "C"}).json()["id"]
        prop = c.post(f"/api/projects/{pid}/proposals", json={"name": "LED-1", "business_line": "LED"}).json()["id"]
        c.put(f"/api/proposals/{prop}/inputs", json={"inputs": SAMPLE["input"]})
        assert c.get("/api/capabilities").json() == {"dwg": True}
        r = c.get(f"/api/proposals/{prop}/drawing.dwg")
        assert r.status_code == 200 and "Drawing_Rev0_DRAFT.dwg" in r.headers["content-disposition"]
        assert b"AVP-CABINET" in r.content
