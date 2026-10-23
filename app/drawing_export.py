"""拼接与线路图的输出（CR-3）：SVG（页面上显示的就是这张）、DXF（分图层、1:1、mm、真正的 DIMENSION 标注）、
DWG（服务器装了 ODA 才有）、箱体清单 Excel、技术方案 PDF（图纸 + 说明栏 + 计算依据）。
所有文件都写入 Rev、规则版本、产品库版本；没批准的 Rev 带 DRAFT。图纸、说明栏和计算依据用同一份数据。"""
import io
import shutil
import subprocess
import tempfile
from html import escape
from pathlib import Path

LAYERS = ["id", "size", "power", "data", "dim", "note"]
CIRCUIT_COLORS = ["#2C5AA0", "#2F7A5B", "#C2410C", "#7A4FA0", "#A0782C", "#1F7A8C", "#8C5A1F", "#D1495B"]
CABLE_COLORS = ["#D1495B", "#2E86AB", "#3B8B5A", "#E59F2B", "#6D597A", "#1F7A8C", "#8C5A1F", "#2C5AA0"]
SIZE_FILL = ["#E3ECF7", "#F6E9D6", "#E6F2EC", "#F1E6F6"]
NAVY, MUTED, RED = "#15203A", "#6B6558", "#A23A2A"
FONT = "-apple-system, 'Segoe UI', 'Microsoft YaHei', 'PingFang SC', sans-serif"
DXF_LAYERS = {"AVP-CABINET": 8, "AVP-CABINET-ID": 7, "AVP-POWER": 5, "AVP-DATA": 1, "AVP-DIM": 3, "AVP-NOTE": 7}

TEXT = {
    "zh": {"title": "拼接与线路图", "project": "项目", "size": "尺寸 · 面积", "cabinets": "箱体", "modules": "模组",
           "resolution": "分辨率", "power": "功率 · 回路", "data": "网线", "reserved": "预留", "spare": "备用",
           "circuit": "回路", "cable": "网线", "box": "箱", "over": "超载", "legend": "图例",
           "numbering": "编号从左下角 R1C1 开始 · 网线起点在左侧",
           "disclaimer": "本图为方案示意，箱体规格以厂家深化为准", "basis": "计算依据", "item": "项目",
           "formula": "公式（代入数值）", "result": "结果", "rule": "规则", "manual": "手动调整"},
    "en": {"title": "Cabinet & cabling drawing", "project": "Project", "size": "Size · Area", "cabinets": "Cabinets",
           "modules": "Modules", "resolution": "Resolution", "power": "Power · Circuits", "data": "Data cables",
           "reserved": "reserved", "spare": "spare", "circuit": "Circuit", "cable": "Cable", "box": "cab.",
           "over": "overload", "legend": "Legend", "numbering": "Numbering from bottom-left R1C1 · data cables start on the left",
           "disclaimer": "Concept drawing; cabinet specs subject to manufacturer detailing", "basis": "Calculation basis",
           "item": "Item", "formula": "Formula (with values)", "result": "Result", "rule": "Rule",
           "manual": "Manually adjusted"},
}


def notes(d: dict, meta: dict, lang: str) -> list[tuple[str, str]]:
    """说明栏（SVG、DXF、PDF 共用）。"""
    t, v = TEXT[lang], meta["values"]
    cabs = " + ".join(f"{k} {s['count']}" for k, s in d["sizes"].items()) + f" = {len(d['cabinets'])}"
    return [
        (t["project"], meta["project"]),
        (t["size"], f"{d['width']} × {d['height']} mm · {v['area_m2']:.2f} ㎡"),
        (t["modules"] + " · " + t["cabinets"], f"{v.get('modules', '—')} · {cabs}"),
        (t["resolution"], f"{v.get('res_x', '—')} × {v.get('res_y', '—')}"),
        (t["power"], f"{v['power_kw']:.2f} kW · {len(d['circuits'])} + {d['spare_circuits']} {t['reserved']}"),
        (t["data"], f"{len(d['cables'])} + {d['spare_cables']} {t['spare']}"),
    ]


def _circled(n: int) -> str:
    return chr(0x2460 + n - 1) if 1 <= n <= 20 else str(n)


def svg(d: dict, meta: dict, lang: str = "zh", layers: list[str] | None = None) -> str:
    """layers：要显示的图层（默认全部）。每个箱体是 <g data-cab="R1C1">，页面上点它可以换回路 / 网线。"""
    t, shown = TEXT[lang], set(LAYERS if layers is None else layers)
    W, H = d["width"], d["height"]
    s = 820 / W
    left, top, gap = 40, 116, 110
    note_x = left + W * s + gap
    width, height = note_x + 330, max(top + H * s + 70, 560)
    X = lambda x: left + x * s  # noqa: E731
    Y = lambda y: top + (H - y) * s  # noqa: E731  CAD 坐标 y 向上
    over_c = {c["no"] for c in d["circuits"] if c["over"]}
    over_d = {c["no"] for c in d["cables"] if c["over"]}
    size_idx = {k: i for i, k in enumerate(d["sizes"])}
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width:.0f} {height:.0f}" width="{width:.0f}"'
           f' height="{height:.0f}" font-family="{FONT}" class="drawing">',
           f'<rect width="100%" height="100%" fill="#FFFFFF"/>',
           f'<text x="{left}" y="26" font-size="16" font-weight="700" fill="{NAVY}">{escape(t["title"])} · {escape(meta["project"])}</text>']
    # 箱体（AVP-CABINET）和图层
    for c in d["cabinets"]:
        x, y, w, h = X(c["x"]), Y(c["y"] + c["h"]), c["w"] * s, c["h"] * s
        stroke = RED if c["circuit"] in over_c or c["cable"] in over_d else (NAVY if c["manual"] else "#8F8676")
        g = [f'<g data-cab="{c["id"]}" class="cab">',
             f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" fill="{SIZE_FILL[size_idx[c["size"]] % 4]}"'
             f' stroke="{stroke}" stroke-width="{2 if stroke != "#8F8676" else 0.8}"/>']
        if "power" in shown:
            col = CIRCUIT_COLORS[(c["circuit"] - 1) % len(CIRCUIT_COLORS)]
            g.append(f'<rect class="layer-power" x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" fill="{col}"'
                     f' fill-opacity="0.14" pointer-events="none"/>'
                     f'<text class="layer-power" x="{x + 3:.1f}" y="{y + h - 4:.1f}" font-size="9.5" font-weight="700"'
                     f' fill="{col}">P{c["circuit"]}</text>')
        if "id" in shown:
            g.append(f'<text class="layer-id" x="{x + w / 2:.1f}" y="{y + h * 0.42:.1f}" font-size="10" text-anchor="middle"'
                     f' fill="{NAVY}">{c["id"]}</text>')
        if "size" in shown:
            g.append(f'<text class="layer-size" x="{x + w / 2:.1f}" y="{y + h * 0.42 + 12:.1f}" font-size="9" text-anchor="middle"'
                     f' fill="{MUTED}">{c["size"]}</text>')
        if "data" in shown:
            col = CABLE_COLORS[(c["cable"] - 1) % len(CABLE_COLORS)]
            g.append(f'<text class="layer-data" x="{x + w - 3:.1f}" y="{y + h - 4:.1f}" font-size="10.5" font-weight="700"'
                     f' text-anchor="end" fill="{col}">{_circled(c["cable"])}</text>')
        g.append("</g>")
        out.append("".join(g))
    if "data" in shown:  # 网线：按蛇形顺序连起来
        by_id = {c["id"]: c for c in d["cabinets"]}
        for cable in d["cables"]:
            col = CABLE_COLORS[(cable["no"] - 1) % len(CABLE_COLORS)]
            pts = " ".join(f"{X(by_id[i]['x'] + by_id[i]['w'] / 2):.1f},{Y(by_id[i]['y'] + by_id[i]['h'] / 2):.1f}"
                           for i in cable["path"])
            out.append(f'<polyline class="layer-data" points="{pts}" fill="none" stroke="{col}" stroke-width="1.6"'
                       f' stroke-opacity="0.8" pointer-events="none"/>')
            if cable["path"]:
                first = by_id[cable["path"][0]]
                out.append(f'<circle class="layer-data" cx="{X(first["x"] + first["w"] / 2):.1f}"'
                           f' cy="{Y(first["y"] + first["h"] / 2):.1f}" r="3.5" fill="{col}" pointer-events="none"/>')
    if "dim" in shown:  # 尺寸线：上方逐列 + 总宽，右侧逐行 + 总高
        y1, y2 = top - 18, top - 46
        for c in range(d["cols"]):
            xa, xb = X(c * d["col_w"]), X((c + 1) * d["col_w"])
            out.append(f'<g class="layer-dim"><line x1="{xa:.1f}" y1="{y1}" x2="{xb:.1f}" y2="{y1}" stroke="{MUTED}"/>'
                       f'<line x1="{xa:.1f}" y1="{y1 - 4}" x2="{xa:.1f}" y2="{y1 + 4}" stroke="{MUTED}"/>'
                       f'<text x="{(xa + xb) / 2:.1f}" y="{y1 - 4}" font-size="9.5" text-anchor="middle" fill="{MUTED}">{d["col_w"]}</text></g>')
        out.append(f'<g class="layer-dim"><line x1="{X(0)}" y1="{y2}" x2="{X(W):.1f}" y2="{y2}" stroke="{NAVY}"/>'
                   f'<text x="{X(W / 2):.1f}" y="{y2 - 5}" font-size="11.5" font-weight="700" text-anchor="middle" fill="{NAVY}">'
                   f'{W} = {d["col_w"]} × {d["cols"]}</text></g>')
        x1, x2, y = X(W) + 16, X(W) + 70, 0
        for h in d["rows"]:
            ya, yb = Y(y), Y(y + h)
            out.append(f'<g class="layer-dim"><line x1="{x1}" y1="{ya:.1f}" x2="{x1}" y2="{yb:.1f}" stroke="{MUTED}"/>'
                       f'<line x1="{x1 - 4}" y1="{yb:.1f}" x2="{x1 + 4}" y2="{yb:.1f}" stroke="{MUTED}"/>'
                       f'<text x="{x1 + 5}" y="{(ya + yb) / 2 + 3:.1f}" font-size="9.5" fill="{MUTED}">{h}</text></g>')
            y += h
        out.append(f'<g class="layer-dim"><line x1="{x2}" y1="{Y(0):.1f}" x2="{x2}" y2="{Y(H):.1f}" stroke="{NAVY}"/>'
                   f'<text x="{x2 + 6}" y="{Y(H / 2):.1f}" font-size="11.5" font-weight="700" fill="{NAVY}"'
                   f' transform="rotate(90 {x2 + 6} {Y(H / 2):.1f})" text-anchor="middle">{H} = {" + ".join(str(h) for h in d["rows"][::-1])}</text></g>')
    if "note" in shown:  # 说明栏
        y = top - 40
        out.append(f'<g class="layer-note"><rect x="{note_x - 12}" y="{y - 22}" width="330" height="{height - y + 10:.0f}"'
                   f' fill="#FAF8F4" stroke="#E6E0D6"/>')
        for label, value in notes(d, meta, lang):
            out.append(f'<text x="{note_x}" y="{y}" font-size="10" fill="{MUTED}">{escape(label)}</text>'
                       f'<text x="{note_x}" y="{y + 14}" font-size="11.5" font-weight="600" fill="{NAVY}">{escape(value)}</text>')
            y += 34
        out.append(f'<text x="{note_x}" y="{y}" font-size="10" fill="{MUTED}">{escape(t["legend"])}</text>')
        y += 16
        for i, (k, sz) in enumerate(d["sizes"].items()):
            out.append(f'<rect x="{note_x}" y="{y - 9}" width="14" height="10" fill="{SIZE_FILL[i % 4]}" stroke="#8F8676"/>'
                       f'<text x="{note_x + 20}" y="{y}" font-size="10.5" fill="{NAVY}">{k} {sz["w"]} × {sz["h"]} × {sz["count"]}</text>')
            y += 16
        for c in d["circuits"]:
            col = RED if c["over"] else CIRCUIT_COLORS[(c["no"] - 1) % len(CIRCUIT_COLORS)]
            tail = f" · {t['over']}" if c["over"] else ""
            out.append(f'<text x="{note_x}" y="{y}" font-size="10.5" fill="{col}">P{c["no"]} {t["circuit"]} {c["no"]} · '
                       f'{c["cabinets"]} {t["box"]} · {c["watts"]} W · {c["amps"]} A{tail}</text>')
            y += 15
        out.append(f'<text x="{note_x}" y="{y}" font-size="10.5" fill="{MUTED}">P{len(d["circuits"]) + 1} {t["circuit"]} '
                   f'{len(d["circuits"]) + 1} · {t["reserved"]}</text>')
        y += 18
        for c in d["cables"]:
            col = RED if c["over"] else CABLE_COLORS[(c["no"] - 1) % len(CABLE_COLORS)]
            tail = f" · {t['over']}" if c["over"] else ""
            out.append(f'<text x="{note_x}" y="{y}" font-size="10.5" fill="{col}">{_circled(c["no"])} {t["cable"]} · '
                       f'{c["cabinets"]} {t["box"]} · {c["px"]:,} px{tail}</text>')
            y += 15
        out.append(f'<text x="{note_x}" y="{y}" font-size="10.5" fill="{MUTED}">{_circled(len(d["cables"]) + 1)} FOR SPARE · {t["spare"]}</text>')
        y += 24
        for line in (t["numbering"], t["disclaimer"], meta["stamp"]):
            out.append(f'<text x="{note_x}" y="{y}" font-size="9.5" fill="{MUTED}">{escape(line)}</text>')
            y += 14
        out.append("</g>")
    if meta["draft"]:
        out.append(f'<text x="{X(W / 2):.1f}" y="{Y(H / 2):.1f}" font-size="{min(W * s / 5, 120):.0f}" font-weight="700"'
                   f' fill="{RED}" fill-opacity="0.10" text-anchor="middle" dominant-baseline="middle"'
                   f' transform="rotate(-20 {X(W / 2):.1f} {Y(H / 2):.1f})" pointer-events="none">DRAFT</text>')
    out.append("</svg>")
    return "".join(out)


def dxf(d: dict, meta: dict, lang: str = "zh") -> bytes:
    """单位 mm，1:1；每样东西一个图层；尺寸用 DIMENSION 实体，在 AutoCAD 里可以直接量、直接改。"""
    import ezdxf

    doc = ezdxf.new("R2018", setup=True)
    doc.header["$INSUNITS"] = 4
    doc.header["$MEASUREMENT"] = 1
    for name, color in DXF_LAYERS.items():
        doc.layers.add(name, color=color)
    msp = doc.modelspace()
    W, H = d["width"], d["height"]
    th = max(W, H) / 120  # 文字高度随屏幕大小
    for c in d["cabinets"]:
        x, y, w, h = c["x"], c["y"], c["w"], c["h"]
        msp.add_lwpolyline([(x, y), (x + w, y), (x + w, y + h), (x, y + h)], close=True, dxfattribs={"layer": "AVP-CABINET"})
        msp.add_text(c["id"], height=th, dxfattribs={"layer": "AVP-CABINET-ID"}).set_placement(
            (x + w / 2, y + h * 0.55), align=ezdxf.enums.TextEntityAlignment.MIDDLE_CENTER)
        msp.add_text(f"{c['size']} {w}x{h}", height=th * 0.7, dxfattribs={"layer": "AVP-CABINET-ID"}).set_placement(
            (x + w / 2, y + h * 0.35), align=ezdxf.enums.TextEntityAlignment.MIDDLE_CENTER)
        msp.add_text(f"P{c['circuit']}", height=th * 0.8, dxfattribs={"layer": "AVP-POWER", "color": (c["circuit"] - 1) % 6 + 1}
                     ).set_placement((x + w * 0.08, y + h * 0.08))
        msp.add_text(f"D{c['cable']}", height=th * 0.8, dxfattribs={"layer": "AVP-DATA", "color": (c["cable"] - 1) % 6 + 1}
                     ).set_placement((x + w * 0.92, y + h * 0.08), align=ezdxf.enums.TextEntityAlignment.BOTTOM_RIGHT)
    by_id = {c["id"]: c for c in d["cabinets"]}
    for cable in d["cables"]:
        pts = [(by_id[i]["x"] + by_id[i]["w"] / 2, by_id[i]["y"] + by_id[i]["h"] / 2) for i in cable["path"]]
        if len(pts) > 1:
            msp.add_lwpolyline(pts, dxfattribs={"layer": "AVP-DATA", "color": (cable["no"] - 1) % 6 + 1})
    dim = {"dimtxt": th, "dimasz": th * 0.8, "dimexe": th * 0.5, "dimexo": th * 0.3, "dimgap": th * 0.3}
    for c in range(d["cols"]):
        msp.add_linear_dim(base=(0, H + th * 4), p1=(c * d["col_w"], H), p2=((c + 1) * d["col_w"], H),
                           override=dim, dxfattribs={"layer": "AVP-DIM"}).render()
    msp.add_linear_dim(base=(0, H + th * 9), p1=(0, H), p2=(W, H), override=dim, dxfattribs={"layer": "AVP-DIM"}).render()
    y = 0
    for h in d["rows"]:
        msp.add_linear_dim(base=(W + th * 4, 0), p1=(W, y), p2=(W, y + h), angle=90, override=dim,
                           dxfattribs={"layer": "AVP-DIM"}).render()
        y += h
    msp.add_linear_dim(base=(W + th * 9, 0), p1=(W, 0), p2=(W, H), angle=90, override=dim,
                       dxfattribs={"layer": "AVP-DIM"}).render()
    t = TEXT[lang]
    lines = [t["title"]] + [f"{a}: {b}" for a, b in notes(d, meta, lang)]
    lines += [f"P{c['no']} {t['circuit']} {c['no']}: {c['cabinets']} {t['box']} · {c['watts']} W · {c['amps']} A"
              + (f" · {t['over']}" if c["over"] else "") for c in d["circuits"]]
    lines += [f"D{c['no']} {t['cable']}: {c['cabinets']} {t['box']} · {c['px']:,} px" + (f" · {t['over']}" if c["over"] else "")
              for c in d["cables"]]
    lines += [f"D{len(d['cables']) + 1} FOR SPARE", t["numbering"], t["disclaimer"], meta["stamp"]]
    msp.add_mtext("\\P".join(lines), dxfattribs={"layer": "AVP-NOTE", "char_height": th, "insert": (W + th * 16, H),
                                                 "width": th * 45})
    if meta["draft"]:
        msp.add_text("DRAFT", height=H / 6, dxfattribs={"layer": "AVP-NOTE", "color": 1}).set_placement(
            (W / 2, H / 2), align=ezdxf.enums.TextEntityAlignment.MIDDLE_CENTER)
    doc.header["$PROJECTNAME"] = meta["project"][:250]
    stream = io.StringIO()
    doc.write(stream)
    return stream.getvalue().encode("utf-8")


def dwg(dxf_bytes: bytes, oda: str) -> bytes:
    """用免费的 ODA File Converter 把 DXF 转成 DWG（ODA 只处理整个文件夹）。"""
    with tempfile.TemporaryDirectory() as src, tempfile.TemporaryDirectory() as dst:
        (Path(src) / "drawing.dxf").write_bytes(dxf_bytes)
        subprocess.run([oda, src, dst, "ACAD2018", "DWG", "0", "1", "*.DXF"], check=True, timeout=180, capture_output=True)
        return next(Path(dst).glob("*.dwg")).read_bytes()


def cabinets_xlsx(d: dict, meta: dict, lang: str = "zh") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    t = TEXT[lang]
    wb = Workbook()
    ws = wb.active
    ws.title = "Cabinets"
    ws.append([f"{t['cabinets']} · {meta['project']}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([meta["stamp"]])
    ws.append([])
    ws.append(["ID", "Size 尺寸", "Circuit 回路", "Data cable 网线", "Row 行", "Column 列", "px", "W", "Manual 手动"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for c in sorted(d["cabinets"], key=lambda c: (c["row"], c["col"])):
        ws.append([c["id"], f"{c['size']} {c['w']} × {c['h']}", c["circuit"], c["cable"], c["row"], c["col"], c["px"],
                   c["watts"], "✓" if c["manual"] else ""])
    for col, width in zip("ABCDEFGHI", (10, 16, 10, 14, 8, 10, 10, 10, 10)):
        ws.column_dimensions[col].width = width
    p = wb.properties
    p.title, p.keywords, p.description, p.creator = f"Cabinets · {meta['project']}", meta["stamp"], meta["stamp"], "AV Proposal Platform"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf(d: dict, meta: dict, basis: list[dict], lang: str = "zh") -> bytes:
    """技术方案 PDF：图纸 + 说明栏 + 计算依据，用无头 Chromium 打印，再用 pypdf 写文件属性。"""
    from playwright.sync_api import sync_playwright
    from pypdf import PdfReader, PdfWriter

    t = TEXT[lang]
    rows = "".join(f"<tr><td>{escape(b['item'])}</td><td class='mono'>{escape(b['formula'])}</td>"
                   f"<td>{escape(b['result'])}</td><td>{escape(b['rule'])}</td></tr>" for b in basis)
    html = f"""<!doctype html><html><head><meta charset="utf-8"><style>
      body {{ font-family: {FONT}; color: {NAVY}; margin: 0; }}
      h1 {{ font-size: 18px; margin: 0 0 4px; }} .muted {{ color: {MUTED}; font-size: 11px; }}
      svg {{ width: 100%; height: auto; }} table {{ border-collapse: collapse; width: 100%; font-size: 11px; margin-top: 10px; }}
      th, td {{ border-bottom: 1px solid #E6E0D6; padding: 5px 6px; text-align: left; vertical-align: top; }}
      .mono {{ font-family: Consolas, monospace; }} h2 {{ font-size: 14px; margin: 14px 0 0; }}
    </style></head><body>
      <h1>{escape(t['title'])} · {escape(meta['project'])}</h1><div class="muted">{escape(meta['stamp'])}</div>
      {svg(d, meta, lang)}
      <h2>{escape(t['basis'])}</h2>
      <table><thead><tr><th>{t['item']}</th><th>{t['formula']}</th><th>{t['result']}</th><th>{t['rule']}</th></tr></thead>
      <tbody>{rows}</tbody></table>
    </body></html>"""
    with sync_playwright() as p:
        browser = p.chromium.launch()
        try:
            page = browser.new_page()
            page.set_content(html)
            raw = page.pdf(format="A3", landscape=True, print_background=True,
                           margin={"top": "12mm", "bottom": "12mm", "left": "12mm", "right": "12mm"})
        finally:
            browser.close()
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(raw)))
    writer.add_metadata({"/Title": f"{t['title']} · {meta['project']}", "/Subject": meta["stamp"],
                         "/Keywords": meta["stamp"], "/Creator": "AV Proposal Platform"})
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def oda_available(oda: str | None) -> bool:
    return bool(oda) and (Path(oda).exists() or shutil.which(oda) is not None)
