"""资料库扫描（CR-8）：资料库文件夹下每类资料一个子文件夹，勾选的类型逐个扫描；新增、改过的文件重读，删掉的文件从索引里拿掉。

- 文件夹：以往项目（一个子文件夹一个项目）、规格书、图纸、BOQ、模板与条款；英文名 projects / specs / drawings / boq / templates 也认。
- 「以往项目」文件夹里的 标准答案.txt / 标准答案.xlsx（每行或第一栏写一个项目文件夹名或项目编号）列出做完、可当标准答案的项目。
- 原文件只读：不移动、不改名、不写入。读到的文字进全文索引 kb_fts，结构化内容存 kb_files.facts。
- 扫完重新整理「学到的内容」：同样的内容（指纹相同）不重复出现，人已经决定过的不再出现。
"""
import hashlib
import logging
import re
import threading
import time
from datetime import datetime, timedelta
from pathlib import Path

from sqlalchemy import delete, select, text
from sqlalchemy.orm import Session

from ..api import iso
from ..models import KbFile, KbItem, KbProject, KbSetting, ProductLibraryVersion, RuleSetVersion, TemplateVersion, utcnow
from ..products.library import load_reference
from ..rules.engine import value_names
from ..versions import latest
from . import learn
from .read import read

TYPES = {"projects": ["以往项目", "projects"], "specs": ["规格书", "specs", "datasheets"], "drawings": ["图纸", "drawings"],
         "boq": ["BOQ"], "templates": ["模板与条款", "templates"]}
STANDARD_FILES = {"标准答案.txt", "标准答案.xlsx", "standard.txt", "standard.xlsx"}
LINES = ["LED", "PJ"]


def type_dirs(root: Path) -> dict[str, Path]:
    existing = {p.name.casefold(): p for p in root.iterdir() if p.is_dir()} if root.is_dir() else {}
    return {t: next((existing[n.casefold()] for n in names if n.casefold() in existing), root / names[0])
            for t, names in TYPES.items()}


def ensure_folders(root: Path) -> None:
    for d in type_dirs(root).values():
        d.mkdir(parents=True, exist_ok=True)


def setting(db: Session, key: str, default):
    row = db.get(KbSetting, key)
    return row.value if row else default


def put_setting(db: Session, key: str, value) -> None:
    row = db.get(KbSetting, key) or KbSetting(key=key)
    row.value, row.updated_at = value, utcnow()
    db.add(row)


def enabled_types(db: Session) -> list[str]:
    return setting(db, "types", {"types": list(TYPES)})["types"]


def lines(db: Session, file_id: int) -> list[tuple[str, str]]:
    return [tuple(r) for r in db.execute(text("SELECT location, text FROM kb_fts WHERE file_id = :f ORDER BY rowid"),
                                         {"f": file_id})]


def _index(db: Session, file_id: int, chunks: list[tuple[str, str]]) -> None:
    db.execute(text("DELETE FROM kb_fts WHERE file_id = :f"), {"f": file_id})
    for location, t in chunks:
        db.execute(text("INSERT INTO kb_fts (text, file_id, location) VALUES (:t, :f, :l)"),
                   {"t": t, "f": file_id, "l": location})


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(1 << 20):
            h.update(chunk)
    return h.hexdigest()


def standard_list(path: Path) -> set[str]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        names = [str(r[0]) for r in wb.active.iter_rows(values_only=True) if r and r[0] is not None]
        wb.close()
    else:
        names = path.read_text(encoding="utf-8-sig", errors="replace").splitlines()
    return {n.strip().casefold() for n in names if n.strip() and not n.strip().startswith("#")}


def collect(root: Path, types: list[str]) -> tuple[dict, set[str]]:
    """{相对路径: (文件, 类型, 项目文件夹名)}，以及标准答案清单里的名字。"""
    found, standard = {}, set()
    for t, d in type_dirs(root).items():
        if t not in types or not d.is_dir():
            continue
        for p in sorted(d.rglob("*")):
            if not p.is_file() or p.name.startswith(("~$", ".")):
                continue
            parts = p.relative_to(d).parts
            if t == "projects" and len(parts) == 1 and p.name.casefold() in STANDARD_FILES:
                standard |= standard_list(p)
                continue
            found[p.relative_to(root).as_posix()] = (p, t, parts[0] if t == "projects" and len(parts) > 1 else None)
    return found, standard


def read_into(db: Session, f: KbFile, path: Path, oda: str | None) -> None:
    r = read(path, oda)
    f.read_method, f.read_note = r["method"], r["note"]
    f.read_status = "failed" if r["method"] == "failed" else "read" if r["chunks"] else "attached"
    f.chunks = len(r["chunks"])
    f.facts = learn.file_facts(r["chunks"], r["rows"], r["method"]) if r["chunks"] else None
    f.read_at = utcnow()
    _index(db, f.id, r["chunks"])


def scan(state, user_id: int | None = None) -> dict | None:
    """扫一遍资料库。同一时间只扫一次：已经在扫时返回 None。"""
    if not state.kb_lock.acquire(blocking=False):
        return None
    try:
        return _scan(state, user_id)
    finally:
        state.kb_progress = None
        state.kb_lock.release()


def _scan(state, user_id: int | None) -> dict:
    root = state.kb_dir
    with state.Session() as db:
        types = enabled_types(db)
    found, standard = collect(root, types)
    summary = {"new": 0, "changed": 0, "removed": 0, "failed": 0, "total": len(found)}
    state.kb_progress = {"done": 0, "total": len(found)}
    with state.Session() as db:
        known = {f.rel_path: f for f in db.scalars(select(KbFile))}
        for rel, (path, t, key) in found.items():
            st, f = path.stat(), known.get(rel)
            if not (f and not f.missing and f.read_status != "queued" and f.size == st.st_size
                    and f.mtime_ns == st.st_mtime_ns):
                digest = sha256(path)
                same = f is not None and not f.missing and f.read_status != "queued" and f.sha256 == digest
                if f is None:
                    f = KbFile(rel_path=rel, doc_type=t, size=st.st_size, mtime_ns=st.st_mtime_ns, sha256=digest)
                    db.add(f)
                    db.flush()
                    summary["new"] += 1
                elif not same:
                    summary["changed" if f.read_status != "queued" else "new"] += 1
                f.size, f.mtime_ns, f.sha256, f.missing = st.st_size, st.st_mtime_ns, digest, False
                if not same:
                    read_into(db, f, path, state.oda)
                    summary["failed"] += f.read_status == "failed"
            f.doc_type, f.project_key = t, key
            db.commit()  # 每读完一份就存，扫描中途停了也不用从头来
            state.kb_progress["done"] += 1
        for rel, f in known.items():
            if rel not in found and not f.missing and f.doc_type in types:
                f.missing = True
                _index(db, f.id, [])
                db.execute(delete(KbItem).where(KbItem.source_file_id == f.id, KbItem.status == "pending"))
                summary["removed"] += 1
        sync_projects(db, standard)
        refresh(db)
        put_setting(db, "last_scan", {"at": iso(utcnow()), "by": user_id, **summary})
        db.commit()
    return summary


def sync_projects(db: Session, standard: set[str]) -> None:
    keys = set(db.scalars(select(KbFile.project_key).where(KbFile.project_key.is_not(None), KbFile.missing.is_(False))))
    have = {p.key for p in db.scalars(select(KbProject))}
    for k in sorted(keys - have):
        m = re.match(r"\s*(\d+)", k)
        db.add(KbProject(key=k, project_no=m[1] if m else None))
    db.flush()
    for p in db.scalars(select(KbProject)):
        p.missing = p.key not in keys
        p.standard_list = p.key.casefold() in standard or bool(p.project_no and p.project_no.casefold() in standard)


def _sync(db: Session, kind: str, fresh: list[dict]) -> None:
    """待确认的旧建议没再出现就删掉；新建议按指纹加入，人决定过的（采用 / 不采用）不再出现。"""
    new = {i["fingerprint"]: i for i in fresh}
    old = {i.fingerprint: i for i in db.scalars(select(KbItem).where(KbItem.kind == kind))}
    for fp, item in old.items():
        if item.status == "pending" and fp not in new:
            db.delete(item)
    for fp, i in new.items():
        if fp not in old:
            db.add(KbItem(**i))


def project_files(db: Session, key: str) -> list[KbFile]:
    return list(db.scalars(select(KbFile).where(KbFile.project_key == key, KbFile.missing.is_(False))
                           .order_by(KbFile.rel_path)))


def refresh(db: Session, kinds: tuple[str, ...] = ("product", "template", "project", "rule")) -> None:
    """按最新的产品库、模板、规则和项目记录重新整理「学到的内容」。"""
    if not kinds:
        return
    files = list(db.scalars(select(KbFile).where(KbFile.missing.is_(False), KbFile.facts.is_not(None))))
    libraries = {line: latest(db, ProductLibraryVersion, line).content for line in LINES}
    if "product" in kinds:
        refs = {line: load_reference(line) for line in LINES}
        _sync(db, "product", [i for f in files if f.doc_type == "specs"
                              if (i := learn.spec_item(f, lines(db, f.id), libraries, refs))])
    if "template" in kinds:
        templates = {line: latest(db, TemplateVersion, line).content for line in LINES}
        known = {line: value_names(latest(db, RuleSetVersion, line).content) for line in LINES}
        boqs = [(f, f.facts["line"], f.facts["boq"]) for f in files if f.facts.get("boq") and f.facts.get("line")]
        _sync(db, "template", learn.template_items(boqs, templates, known))
    projects = list(db.scalars(select(KbProject).where(KbProject.missing.is_(False)).order_by(KbProject.key)))
    if "project" in kinds:
        projectors = [p for p in [*libraries["PJ"]["products"], *load_reference("PJ")] if p["category"] == "projector"]
        items = []
        for p in projects:
            rec = learn.project_record([{"id": f.id, "facts": f.facts, "lines": lines(db, f.id)}
                                        for f in project_files(db, p.key) if f.facts], projectors)
            if rec and rec != p.record:
                first = next((s for s in rec["sources"].values() if s["file_id"]), {})
                items.append({"kind": "project", "business_line": rec["line"],
                              "fingerprint": learn.fingerprint("project", p.key, rec),
                              "payload": {"project_key": p.key, "record": rec},
                              "source_file_id": first.get("file_id"), "location": first.get("location", ""), "snippet": "",
                              "confidence": rec["confidence"]})
        _sync(db, "project", items)
    if "rule" in kinds:
        items = []
        for p in projects:
            if p.record and (p.standard_manual or p.standard_list):
                rules = latest(db, RuleSetVersion, p.record["line"])
                result = learn.compare(p.record, rules.content, libraries[p.record["line"]])
                items += learn.rule_items(p.key, p.record, result, rules.version)
        _sync(db, "rule", items)
    db.flush()


def nightly(state) -> threading.Thread:
    """每晚在 knowledge_scan_at（服务器本地时间）复扫一次；正在扫时跳过这次。"""
    def loop():
        while True:
            now = datetime.now()
            h, m = map(int, state.kb_scan_at.split(":"))
            at = now.replace(hour=h, minute=m, second=0, microsecond=0)
            time.sleep(((at if at > now else at + timedelta(days=1)) - now).total_seconds())
            try:
                scan(state)
            except Exception:  # 复扫失败不能让服务停下；下一晚再试
                logging.exception("资料库夜间复扫失败")

    t = threading.Thread(target=loop, daemon=True, name="kb-nightly")
    t.start()
    return t
