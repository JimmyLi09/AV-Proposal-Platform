"""CR-6 验收：去掉价格，客户版 BOQ（143-Royal Plaza 出 7 行），描述 / 规格 / 数量可改，改数量要写原因，可导出 Excel。"""
import copy
import io
import json
import re
from pathlib import Path

import pytest
import sqlalchemy as sa
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.boq import build
from app.db import MIGRATIONS, make_engine, migrate
from app.main import create_app
from app.products.library import load_seed as load_library
from app.rules.engine import calculate, load_seed, value_names
from app.templates.content import TemplateError, load_seed as load_template, validate

RULES, SAMPLE = load_seed("LED")
LIBRARY = load_library("LED")
TEMPLATE = load_template("LED")
GOLDEN = json.loads((Path(__file__).parent / "fixtures" / "regression_cases.json").read_text(encoding="utf-8"))
CASE = GOLDEN["led_royal_plaza_boq"]
PRICE_WORDS = re.compile(r"price|cost|gst|subtotal|markup|sell", re.I)


def boq_143(edits=None):
    out = calculate(RULES, LIBRARY, SAMPLE["input"])
    return build(TEMPLATE, out["values"] | {"proposal_name": "LED-1 Lobby"}, edits)


def test_143_client_boq_has_the_7_golden_lines():
    boq = boq_143()
    assert [[r["title"], r["qty"], r["unit"]] for r in boq["rows"]] == CASE["expect_lines"]
    led = boq["rows"][0]
    assert [s.split(":")[0] for s in led["specs"]] == CASE["led_spec_rows"]
    assert led["specs"] == ["Overall dimension: 4480 × 2560 mm",
                            "Cabinet size: 640 × 480 / 640 × 640 (module 320 × 160)",
                            "Brightness: 800 cd/㎡ (≥ 600 required), adjustable",
                            "Overall screen resolution: 1792 × 1024", "Aspect ratio: 7:4",
                            "Spares: 12 modules, 1 power supply, 1 receiving card"]
    assert boq["rows"][3]["note"] == "Client may provide"
    assert boq["rows"][4]["specs"] == ["DB box with MCB × 3 circuits",
                                       "Power cables 3 + 1 spare, data cables CAT6 4 + 1 spare"]
    assert boq["group"] == "LED-1 Lobby" and boq["manual"] == []


def test_seed_template_is_valid_and_rejects_bad_rows():
    known = value_names(RULES)
    validate(TEMPLATE, "LED", known)
    for change, msg in [(lambda t: t["boq"]["rows"][0].update(title="{nope} screen"), "不认识"),
                        (lambda t: t["boq"]["rows"][0].update(unit="m"), "单位"),
                        (lambda t: t["boq"]["rows"][1].update(qty="price_total"), "数量"),
                        (lambda t: t["boq"]["rows"][1].update(qty=-1), "数量"),
                        (lambda t: t["boq"]["rows"][1].update(key="led_screen"), "不能重复"),
                        (lambda t: t["boq"]["rows"][2]["specs"].append("{controller.__class__}"), "花括号")]:
        bad = copy.deepcopy(TEMPLATE)
        change(bad)
        with pytest.raises(TemplateError, match=msg):
            validate(bad, "LED", known)


def test_rule_set_and_library_have_no_price_fields():
    assert not [p["name"] for p in RULES["params"] if PRICE_WORDS.search(p["name"] + p["group"])]
    assert not PRICE_WORDS.search(json.dumps(LIBRARY))


# ---- API ----

@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path)) as c:
        user = c.post("/api/users", json={"name": "Jimmy", "role": "PM"}).json()
        c.post("/api/session", json={"user_id": user["id"]})
        pid = c.post("/api/projects", json={"name": "143-Royal Plaza LED", "client": "Royal Plaza"}).json()["id"]
        c.proposal = c.post(f"/api/projects/{pid}/proposals", json={"name": "LED-1 Lobby", "business_line": "LED"}).json()["id"]
        r = c.put(f"/api/proposals/{c.proposal}/inputs", json={"inputs": SAMPLE["input"]})
        assert r.json()["rev"]["status"] == "calculated"
        yield c


def keys(obj):
    if isinstance(obj, dict):
        for k, v in obj.items():
            yield k
            yield from keys(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from keys(v)


def test_no_prices_in_any_api(client):
    view = client.get(f"/api/proposals/{client.proposal}").json()
    assert not [k for k in keys(view) if PRICE_WORDS.search(k)]
    assert client.get("/api/costs").status_code in (404, 405)
    assert not [p for p in client.get("/api/rules/LED").json()["content"]["params"] if p["name"] == "gst_rate"]


def save(client, body):
    return client.put(f"/api/proposals/{client.proposal}/boq", json=body)


def test_edit_qty_needs_reason_then_shows_manual_and_restores(client):
    r = save(client, {"rows": {"power_cabling": {"qty": 2}}})
    assert r.status_code == 422 and "原因" in r.json()["detail"]

    r = save(client, {"group": "Lobby LED Wall · LED-1",
                      "rows": {"power_cabling": {"qty": 2, "reason": "两个配电点，分开计"},
                               "led_screen": {"title": "P2.5 Indoor LED wall 4480 × 2560 mm",
                                              "specs": ["Overall dimension: 4480 × 2560 mm", "  "]}}})
    assert r.status_code == 200, r.text
    boq = r.json()["boq"]
    assert boq["group"] == "Lobby LED Wall · LED-1"
    led, power = boq["rows"][0], boq["rows"][4]
    assert led["title"] == "P2.5 Indoor LED wall 4480 × 2560 mm" and led["specs"] == ["Overall dimension: 4480 × 2560 mm"]
    assert not led["manual"]  # 改描述 / 规格不算改数量
    assert (power["qty"], power["rule_qty"], power["manual"], power["reason"], power["by"]) == (2, 1, True, "两个配电点，分开计", "Jimmy")
    assert [(m["key"], m["rule_qty"], m["qty"]) for m in boq["manual"]] == [("power_cabling", 1, 2)]

    boq = save(client, {"rows": {"power_cabling": {"qty": 1}}}).json()["boq"]  # 改回规则值 = 恢复
    assert boq["rows"][4]["manual"] is False and boq["manual"] == [] and boq["group"] == "LED-1 Lobby"
    assert boq["rows"][0]["title"].startswith("P2.5 Indoor LED screen")
    assert save(client, {"rows": {"nope": {"qty": 1}}}).status_code == 422


def test_add_line_is_numbered_and_listed_as_manual(client):
    boq = save(client, {"added": [{"title": "Content player", "specs": ["4K, 1 output"], "qty": 1, "unit": "set"}]}).json()["boq"]
    row = boq["rows"][-1]
    assert (row["no"], row["key"], row["added"], row["manual"], row["by"]) == (8, "x1", True, True, "Jimmy")
    boq = save(client, {"added": [{"key": "x1", "title": "Content player", "specs": ["4K, 1 output"], "qty": 1, "unit": "set"},
                                  {"title": "Signage CMS licence", "qty": 1, "unit": "lot"}]}).json()["boq"]
    assert [r["key"] for r in boq["rows"][7:]] == ["x1", "x2"] and len(boq["manual"]) == 2
    assert save(client, {"added": [{"title": "x", "qty": 1, "unit": "m"}]}).status_code == 422


def test_boq_excel_export(client):
    save(client, {"rows": {"power_cabling": {"qty": 2, "reason": "两个配电点"}}})
    r = client.get(f"/api/proposals/{client.proposal}/boq.xlsx")
    assert r.status_code == 200 and "BOQ_Rev0.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    rows = list(wb["BOQ"].iter_rows(values_only=True))
    assert rows[2][0] == "DRAFT · Rev 0 · Rules v1.3 · Product library v1.0 · Template v1.0"
    assert rows[4] == ("No.", "Description", "Qty", "Unit")
    assert rows[5][1] == "LED-1 Lobby"
    items = [r for r in rows[6:] if r[0]]
    assert [(r[0], r[1], r[2], r[3]) for r in items] == [
        (i + 1, t, 2 if i == 4 else q, u) for i, (t, q, u) in enumerate(CASE["expect_lines"])]
    assert rows[7][1] == "- Overall dimension: 4480 × 2560 mm"
    assert any(r[1] == "*Client may provide" and r[0] is None for r in rows)
    assert wb.properties.keywords == rows[2][0] and wb.properties.creator == "AV Proposal Platform"
    assert not any(PRICE_WORDS.search(str(c)) for r in rows for c in r if c)


def test_template_edit_makes_new_version_used_by_boq(client):
    t = client.get("/api/templates/LED").json()
    assert t["version"] == "v1.0" and "area_m2" in t["names"] and t["units"] == ["sqm", "set", "lot", "pcs"]
    content = copy.deepcopy(t["content"])
    content["boq"]["rows"][3]["title"] = "Media player PC"
    r = client.post("/api/templates/LED", json={"base_id": t["id"], "reason": "改成播放电脑", "content": content})
    assert r.status_code == 201 and r.json()["version"] == "v1.1"
    v = client.get(f"/api/proposals/{client.proposal}").json()
    assert v["template_version"] == "v1.1" and v["boq"]["rows"][3]["title"] == "Media player PC"

    assert client.post("/api/templates/LED", json={"base_id": t["id"], "reason": "x", "content": content}).status_code == 409
    bad = copy.deepcopy(content)
    bad["boq"]["rows"][0]["qty"] = "unknown_value"
    r = client.post("/api/templates/LED", json={"base_id": r.json()["id"], "reason": "x", "content": bad})
    assert r.status_code == 422
    assert [x["version"] for x in client.get("/api/templates/LED/versions").json()] == ["v1.1", "v1.0"]


def test_upgrade_removes_gst_with_a_new_rule_version_and_drops_costs(tmp_path):
    engine = make_engine(tmp_path / "avp.db")
    cfg = Config()
    cfg.set_main_option("script_location", str(MIGRATIONS))
    with engine.begin() as conn:
        cfg.attributes["connection"] = conn
        command.upgrade(cfg, "0005")
    old = copy.deepcopy(RULES)
    old["params"].append({"name": "gst_rate", "group": "Commercial 商务", "label": "GST", "unit": "", "value": 0.09,
                          "status": "sample"})
    sets = sa.table("rule_set_versions", sa.column("business_line"), sa.column("version"),
                    sa.column("content", sa.JSON), sa.column("reason"), sa.column("created_at", sa.DateTime))
    with engine.begin() as conn:
        conn.execute(sets.insert().values(business_line="LED", version="v1.4", content=old, reason="试用",
                                          created_at=sa.func.now()))
    migrate(engine)
    with engine.connect() as conn:
        rows = conn.execute(sa.text("select version, content, reason from rule_set_versions order by id")).all()
        tables = set(sa.inspect(conn).get_table_names())
    assert [r.version for r in rows] == ["v1.4", "v1.5"]
    assert "gst_rate" in rows[0].content and "gst_rate" not in rows[1].content  # 旧版本不改
    assert json.loads(rows[1].content)["rules"] == old["rules"] and "CR-6" in rows[1].reason
    assert "cost_entries" not in tables and "template_versions" in tables
