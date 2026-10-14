"""生成第 4 步的测试样本（自制，非客户资料）。真实样本（清单 G1、G2）到了以后放到同一个目录旁边对比。

用法：python tests/fixtures/samples/make_samples.py   （需要 requirements-dev.txt 里的 reportlab）
内容仿照原型里的 143-Royal Plaza LED-1。
"""
import random
from pathlib import Path

import ezdxf
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import A3, landscape
from reportlab.pdfgen import canvas

OUT = Path(__file__).parent
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def text_pdf():
    """带文字层的立面图 PDF：第 2 页标了屏幕尺寸和点间距。"""
    c = canvas.Canvas(str(OUT / "Lobby_Elevation_RevB.pdf"), pagesize=landscape(A3))
    c.setFont("Helvetica-Bold", 20)
    c.drawString(60, 760, "LOBBY ELEVATION A")
    c.setFont("Helvetica", 12)
    c.drawString(60, 730, "FFL +0.000   CL +3.600")
    c.rect(300, 300, 500, 280)
    c.showPage()
    c.setFont("Helvetica", 14)
    c.drawString(60, 760, "LED SCREEN 4500 (W) x 2560 (H)")
    c.drawString(60, 735, "BOTTOM OF SCREEN +0.900")
    c.drawString(60, 710, "INDOOR LED WALL, PIXEL PITCH P2.5")
    c.rect(300, 300, 450, 256)
    c.save()


def excel():
    """需求表：栏名在左、值在右。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "LED"
    ws["A1"] = "LED Requirement"
    for row, (label, value) in {4: ("Indoor / Outdoor", "Indoor"), 6: ("Pixel pitch", "P2.5"),
                                9: ("Installation", "Wall mount"), 10: ("Maintenance", "Front"),
                                11: ("Brightness (nits)", 600)}.items():
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
    wb.save(OUT / "LED_Requirement.xlsx")


def _image(lines, size=(1400, 700), bg="white", angle=0.0, noise=False):
    img = Image.new("RGB", size, bg)
    d = ImageDraw.Draw(img)
    for y, (text, px) in enumerate(lines):
        d.text((60, 80 + y * 120), text, fill=(30, 30, 30), font=ImageFont.truetype(FONT, px))
    if angle:
        img = img.rotate(angle, expand=True, fillcolor=bg)
    if noise:
        rnd = random.Random(1)
        px = img.load()
        for _ in range(img.width * img.height // 40):
            x, y = rnd.randrange(img.width), rnd.randrange(img.height)
            px[x, y] = (rnd.randrange(150, 230),) * 3
        img = img.filter(ImageFilter.GaussianBlur(0.6))
    return img


def photo():
    """手机拍的现场照片：有点歪、有噪点。"""
    _image([("SCREEN W4480 X H2560", 44), ("VIEWING DISTANCE 4-6M", 44)], bg=(225, 222, 214), angle=2.0,
           noise=True).save(OUT / "Lobby_site_photo.jpg", quality=85)


def scanned_pdf():
    """扫描版平面图：PDF 里只有图片，没有文字层。"""
    img = _image([("PLAN VIEW", 40), ("LED 4480 x 2560", 40)], size=(1600, 1100))
    png = OUT / "_scan.png"
    img.save(png)
    c = canvas.Canvas(str(OUT / "Scanned_Plan.pdf"), pagesize=landscape(A3))
    c.drawImage(str(png), 0, 0, width=landscape(A3)[0], height=landscape(A3)[1])
    c.save()
    png.unlink()


def render():
    """概念效果图：没有尺寸。"""
    img = Image.new("RGB", (1200, 700), (40, 50, 70))
    d = ImageDraw.Draw(img)
    d.rectangle((300, 200, 900, 500), fill=(90, 140, 200))
    d.text((60, 60), "CONCEPT", fill="white", font=ImageFont.truetype(FONT, 40))
    img.save(OUT / "Concept_render.png")


def dxf():
    """CAD 立面（DXF）：标注写在文字里。"""
    doc = ezdxf.new()
    msp = doc.modelspace()
    msp.add_text("LED SCREEN W4480 x H2560", dxfattribs={"layer": "A-ANNO", "height": 100}).set_placement((0, 3000))
    msp.add_mtext("PIXEL PITCH P2.5\\PINDOOR", dxfattribs={"layer": "A-ANNO", "char_height": 80, "insert": (0, 2700)})
    msp.add_lwpolyline([(0, 0), (4480, 0), (4480, 2560), (0, 2560)], close=True)
    doc.saveas(OUT / "Elevation.dxf")


if __name__ == "__main__":
    for make in (text_pdf, excel, photo, scanned_pdf, render, dxf):
        make()
    print("samples written to", OUT)
